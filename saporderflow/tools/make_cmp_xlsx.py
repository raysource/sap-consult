#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成 SAP受注形態比較_講義テキスト.xlsx
-- SAP 受注形態の横断比較サイト（saporderflow/）の Excel 版 --

用法: python3 tools/make_cmp_xlsx.py
依赖: openpyxl

シート構成:
  0_概要 / 1_比較マトリクス / 2_早見表 / 3_E vs Q 対照 / 4_演習解答 / 5_出典
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.abspath(os.path.join(HERE, "..", "SAP受注形態比較_講義テキスト.xlsx"))

BLUE, GREY, INK = "0A6ED1", "F2F4F7", "1D2D3E"
F_TITLE = Font(name="微软雅黑", size=15, bold=True, color="FFFFFF")
F_SMALL = Font(name="微软雅黑", size=9, color="6B7A8D")
F_HEAD = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
F_BODY = Font(name="微软雅黑", size=10, color=INK)
FILL_TITLE = PatternFill("solid", fgColor=BLUE)
FILL_HEAD = PatternFill("solid", fgColor=BLUE)
FILL_ALT = PatternFill("solid", fgColor=GREY)
THIN = Side(style="thin", color="C9D6E0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")


def table(ws, title, subtitle, headers, rows, widths):
    ncol = len(headers); last = get_column_letter(ncol)
    ws.merge_cells(f"A1:{last}1")
    c = ws.cell(1, 1, title); c.font = F_TITLE; c.fill = FILL_TITLE
    ws.row_dimensions[1].height = 26
    ws.merge_cells(f"A2:{last}2"); ws.cell(2, 1, subtitle).font = F_SMALL
    ws.row_dimensions[2].height = 15
    for j, h in enumerate(headers, 1):
        c = ws.cell(3, j, h); c.font = F_HEAD; c.fill = FILL_HEAD; c.alignment = CENTER; c.border = BORDER
    ws.row_dimensions[3].height = 22
    for i, row in enumerate(rows):
        r = 4 + i
        for j, v in enumerate(row, 1):
            c = ws.cell(r, j, v); c.font = F_BODY; c.alignment = WRAP; c.border = BORDER
            if i % 2 == 1:
                c.fill = FILL_ALT
        ws.row_dimensions[r].height = max(17, 13 * (1 + max((len(str(x)) for x in row), default=0) // 30))
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(4, 1)


def kv(ws, title, subtitle, pairs, widths):
    ncol = len(widths); last = get_column_letter(ncol)
    ws.merge_cells(f"A1:{last}1")
    c = ws.cell(1, 1, title); c.font = F_TITLE; c.fill = FILL_TITLE
    ws.row_dimensions[1].height = 26
    ws.merge_cells(f"A2:{last}2"); ws.cell(2, 1, subtitle).font = F_SMALL
    for i, (k, v) in enumerate(pairs):
        r = 3 + i
        ck = ws.cell(r, 1, k); cv = ws.cell(r, 2, v)
        ck.font = F_HEAD; ck.fill = FILL_HEAD; ck.alignment = WRAP; ck.border = BORDER
        cv.font = F_BODY; cv.alignment = WRAP; cv.border = BORDER
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=ncol)
        ws.row_dimensions[r].height = max(17, 13 * (1 + len(str(v)) // 60))
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


SUMMARY = [
    ("教材名称", "SAP 受注形態の横断比較（MTS / MTO / ETO / ATO / VC）"),
    ("教材ID", "TRN-ORDERFORMS-2026"),
    ("目的", "「どの生産形態を選ぶか」で後工程（库存区分・所要量タイプ・成本対象・請求方式・収益認識）がどう連鎖するかを 1 枚で比較し、E（受注在庫）と Q（プロジェクト在庫）の対比を通じて説明力を身につける。"),
    ("対象範囲", "5 形態の比較マトリクス／選型判断／E vs Q の STEP 対照／VC（バリアント設定）の概観／演習とテスト。"),
    ("対象外", "実機での詳細手顺（姉妹サイト: sapmto＝MTO、sapeto＝ETO を参照）、VC の作り込み（特徴設計・依存関係の実装）、EPPM/ポートフォリオ。"),
    ("システム前提", "SAP S/4HANA On-Premise 2020〜2024（画面基準 2023/2024）。Cloud はスコープアイテム（6GD / IYT / 7DM 等）の Test Script を優先。"),
    ("関連教材", "~/Desktop/work/training/sapmto/（MTO：E・受注在庫・戦略 20）、~/Desktop/work/training/sapeto/（ETO：Q・プロジェクト在庫・WBS）"),
    ("完成基準", "テスト 20 題で 75% 以上（15 問）＋「需要の帰属・在庫の表・成本の対象・請求の仕方・月末処理」の 5 点を資料なしで 5 分で説明できること。"),
    ("本教材の立場", "比較表の値は「標準の通常情況」。ETO（戦略グループ・WBS 割当可能な明細カテゴリ）と VC（モデリング設計）は環境差・設計差が大きいため、具体値は自システムで実測する。"),
    ("改訂", "v1.0（2026-09）初版。サイト版（saporderflow/ の 6 ページ）と内容同期。"),
]

MATRIX = [
    ("1 計画起点", "PIR（MD61）", "受注", "受注 ＋ プロジェクト（WBS）", "受注（保存と同時に指図）", "受注 ＋ 設定値（起点は組合せ次第）"),
    ("2 库存区分", "自由在庫", "E 受注在庫", "Q プロジェクト在庫", "E（または自由在庫＋指図）", "E または Q"),
    ("3 需求の単位", "品目＋プラント（予測）", "受注番号＋明細", "WBS 要素", "受注番号＋明細", "受注明細＋設定値（＋WBS）"),
    ("4 調達/製造", "MRP → 計画手配 → 指図", "MRP → 計画手配 → 指図", "プロジェクト/受注 MRP（MD51/MD50）→ 指図・ネットワーク・購買", "受注保存時に指図を自動生成（戦略 82）", "設定値で BOM/作業手順を選別 → 指図"),
    ("5 成本の対象", "品目/指図", "製造指図（受注明細へ決済）", "WBS 要素（＋指図）", "製造指図", "指図（＋WBS の場合あり）"),
    ("6 請求方式", "出荷ベース（F2）", "出荷ベース（F2）", "請求計画・出来高／出荷ベース／RRB（DP90/DP91）", "出荷ベース", "出荷ベース（変種価格の加算を含む）"),
    ("7 収益認識", "出荷時点", "出荷時点", "進捗 → 結果分析（KKA2）→ WIP（KKAX）→ 決済", "出荷時点", "出荷時点（VC 固有ではない）"),
    ("8 主数据の分かれ目", "MRP タイプ・PIR", "MRP4 個別所要量＋明細カテゴリの特別在庫 E＋戦略 20", "MRP4 個別所要量＋評価付プロジェクト在庫＋戦略（Cloud E2）＋受注明細の WBS", "戦略 82（所要量クラス 201・指図タイプ PP04）＋BOM", "品目タイプ KMAT＋品目カテゴリグループ 0002＋特徴/クラス/設定プロファイル"),
    ("9 主要 T-code", "MD61・MD01N・CO01・VL01N・VF01", "VA01・MD04・CO08・VL02N・VF01", "CJ20N・CJ40/CJ42・CJ30/CJ32・MD51・CJI3・KKA2・CJ88", "VA01（保存＝指図作成）・COOIS", "CT04・CL02・CU41/PMEVC・CU50・VA01（構成）"),
]

QUICK1 = [
    ("10 / 11 / 30 / 40", "見込生産系", "（独立所要量系）", "—", "PIR を使う。MTO では使わない"),
    ("20", "受注生産（組立指図なし）", "KE", "040 系", "MRP 経由で指図を作る標準的な MTO"),
    ("25", "受注生産（設定可能品目）", "KEK", "046", "VC × MTO の組合せ"),
    ("82", "組立処理（受注と同時に指図）", "KMFA", "201（指図タイプ PP04）", "受注と指図は 1:1・BOM 必須"),
    ("E2", "受注設計生産（Cloud の ETO）", "E21", "—", "特別在庫 Q。スコープアイテム 6GD の活性後に利用可能（KBA 3623492）"),
]
QUICK2 = [
    ("（空）", "自由在庫（プラント在庫）", "MARD / MCHB（バッチ）", "プラント / 保管場所", "どの受注でも引当可"),
    ("E", "受注在庫（得意先個別在庫）", "MSKA", "受注番号 ＋ 明細", "その受注明細の出荷で引当"),
    ("Q", "プロジェクト在庫", "MSPR", "WBS 要素", "その WBS を指す出荷で引当"),
    ("K / O / W", "預託（仕入先 / 外注支給 / 得意先）", "MKOL / MSLB / MSKU", "仕入先・外注先・得意先", "外注・預託の文脈"),
]
QUICK3 = [
    ("0001 / NORM", "標準的な在庫品（MTS・MTO の完成品）", "TAN（標準）", "VOV4 の決定キー"),
    ("0002 / LUMF", "VC（設定可能品目）および BOM 関連", "—（設定が入る）", "VC の受注で BOM 展開・設定入力を行うためのキー"),
    ("0004", "VC の SET 処理など", "—", "同上（バリエーションによる）"),
    ("ERLA", "設定可能製品のヘッダ（親）", "—", "受注 BOM の親側"),
    ("DIEN / LEIS", "サービス品目", "納入不可／納入関連のサービス明細", "サービス契約・請求のみの明細"),
]

EQ = [
    ("0 品目マスタ", "FERT・MRP1 PD/EX/E・MRP4 個別所要量・戦略 20", "同じ 3 点 ＋ 戦略グループ（Q 系）・会計ビューの評価クラス", "同じ土台。ETO は「Q を指す戦略」が加わる"),
    ("1 事前準備", "（不要）", "CJ20N でプロジェクト＋WBS（請求要素/勘定設定要素）＋評価付在庫 ON ＋ネットワーク", "完全に違う：ETO だけが案件の器を先に作る"),
    ("2 受注", "VA01：得意先・品目・数量・納期のみ", "VA01：＋明細に WBS 要素を入力（＋請求計画）", "違う：ETO は需要の帰属先（WBS）を人が指定"),
    ("3 需要の確認", "MD04：得意先所要量（受注番号＋明細付き）", "MD04：得意先所要量（WBS 付き）", "同じ画面・違う帰属"),
    ("4 MRP 実行", "MD01N / MD02（単品目）", "MD51（プロジェクト）/ MD50（受注）/ MD01N", "入口が違う。結果は同じく MD04 で見る"),
    ("5 計画手配", "在庫区分 = 受注在庫 E、受注番号付き", "勘定設定 = WBS、在庫区分 = プロジェクト在庫 Q", "ここが分岐点。以降の入出庫・評価・成本が従う"),
    ("6 製造指図", "CO08（受注紐づき）/ 計画手配から変換。受注明細参照", "CO40/CO08。勘定設定 = WBS（AUFK-PS_PSP_PNR）", "指図は同じ道具。誰のものかが違う"),
    ("7 部品出庫・実績", "MIGO 261 ／ CO11N", "同じ（＋ネットワークなら CN25 で活動確認）", "ほぼ同じ。ETO は「活動」の受け皿が増える"),
    ("8 完成品入庫", "101 → MSKA（受注在庫 E）", "101 → MSPR（プロジェクト在庫 Q）", "違う表に入る。画面は両方「特別在庫」欄に出るため表で確認"),
    ("9 出荷", "VL01N（在庫区分 E）→ PGI 601", "VL01N（在庫区分 Q ＋ WBS）→ PGI 601（または CNSO）", "操作は同じ・引当先が違う"),
    ("10 請求", "VF01（出荷ベース F2）", "VF01（請求計画/出来高が主。着手金 30% ＋ 完成 70%）", "收钱的方式不同"),
    ("11 月末処理", "差異計算 KKS1 → 指図決済 KO88（受注明細へ）", "進捗 CNE1 → 結果分析 KKA2/KKAJ → WIP KKAX → 決済 CJ88（WBS）", "一番差が出る：指図の損益 vs 案件の進捗損益"),
    ("12 レポート", "MD04・MMBE・VBBS・指図の原価分析", "CJI3・CJ31・CJIB・CJIC", "見る帳票が違う"),
]

SOLUTIONS = [
    ("振り分け-①", "標準的な工作機械の受注生産（納期 2 か月・分割請求なし）", "MTO（E）", "在庫 E・成本=指図・請求=出荷ベース"),
    ("振り分け-②", "得意先工場に据え付ける特注ライン（設計・据付・検収・着手金）", "ETO（Q）", "在庫 Q・成本=WBS・請求=請求計画/出来高＋結果分析"),
    ("振り分け-③", "自社ブランドの切削工具（見込みで在庫し出荷）", "MTS", "在庫 自由・計画 PIR・請求=出荷ベース"),
    ("振り分け-④", "得意先が容量・色・材質を選ぶ産業用ポンプ（受注後すぐ組立）", "VC ＋ ATO", "KMAT+0002・戦略 82・在庫 E または自由＋指図"),
    ("振り分け-⑤", "得意先ごとにソフトウェアをカスタマイズして納める案件", "ETO（サービス明細の併用も）", "在庫 Q または在庫なし＋費用・請求計画/出来高"),
    ("判別-1", "受注 10 台・納期 2 か月・分割請求なし・設計は流用", "MTO（E）", "WBS を作る理由がない。原価は指図で足りる"),
    ("判別-2", "工場据付の特注ライン・着手金 30%・検収後 70%", "ETO（Q）", "案件単位の原価/収益管理 → WBS＋請求計画"),
    ("判別-3", "MMBE に「特別在庫 E、受注 0000001234/000010」", "MTO", "特別在庫 E は受注明細に帰属。表は MSKA"),
    ("判別-4", "計画手配の勘定設定が空・在庫区分も空。最初に疑う設定は？", "4 点セットの不成立", "受注明細の帰属（E の特別在庫 / Q の WBS）→ 品目 MRP4 → 戦略グループの順"),
    ("判別-5", "CJI3 に収益が出ない", "受注明細が指す WBS が請求要素でない", "収益は請求要素、原価は勘定設定要素"),
    ("判別-6", "請求計画なしで出荷ベースだけの請求は可能か？", "可能", "請求計画は分割請求の仕組み。明細カテゴリの請求関連と VTFL を確認"),
    ("判別-7", "同じ品目を MTO と ETO の両方で使えるか？", "技術的には可能（運用は分けるのが安全）", "E と Q が混在すると入庫先・評価・決済が読みにくい"),
    ("判別-8", "得意先が色と容量を選ぶ。MTO のままで対応できるか？", "不可（VC が必要）", "KMAT・特徴/クラス・設定プロファイルと依存関係。品目カテゴリグループ 0002＋戦略 25 または 82"),
]

SOURCES = [
    ("SAP Help（S/4HANA On-Premise）", "Special Settings for Production Orders", "戦略 82＝組立処理（KMFA / 201 / PP04、受注と指図 1:1）", "help.sap.com/docs/SAP_S4HANA_ON-PREMISE/eedc1019283a438a8b73fdde490abc4f/d242b853ff98b44ce10000000a174cb4.html"),
    ("SAP Help（SUPPORT_CONTENT）", "Stock Tables and Stock Types", "MSKA＝受注在庫(E)、MSPR＝プロジェクト在庫(Q)、MSKU＝預託(V/W)", "help.sap.com/docs/SUPPORT_CONTENT/erpscm/3362167795.html"),
    ("SAP Help（Cloud）", "受注設計生産 (ETO) / Make-to-Order (MTO)", "個別所要量品目、評価付販売伝票在庫（所要量クラス 046・評価 M）", "help.sap.com/docs/SAP_S4HANA_CLOUD/4032610758dc437089f0c28320eec93f/af08789e7a6842f68acb056a0bfede9a.html"),
    ("SAP Help（Cloud）", "販売伝票の企業ポートフォリオおよびプロジェクト管理", "受注明細に WBS を割り当てられる明細カテゴリ（例 CBAO）", "help.sap.com/docs/SAP_S4HANA_CLOUD/a376cd9ea00d476b96f18dea1247e6a5/3804ed4c209a408d8eb660a67172139a.html"),
    ("SAP Help / SAP PRESS", "Advanced Variant Configuration / What Is SAP AVC and How Does It Compare to LO-VC", "S/4HANA の AVC（PMEVC）と従来 LO-VC の関係", "help.sap.com/docs/SAP_S4HANA_ON-PREMISE/.../76276945f5d24a119ff4aeaca6ded355.html"),
    ("varconf.com", "Make-to-Order with SAP VC / VC Object Dependencies / VC Pricing", "品目カテゴリグループ 0002・0004、戦略 25→KEK→046、依存関係の種類、変種価格", "varconf.com"),
    ("SAP Community", "VC 品目マスタの必須設定／変種価格（SDCOM-VKOND と VA00）", "KMAT・0002/0004・戦略 25・MRP タイプ PD・個別/包括所要量", "community.sap.com"),
    ("SAP Knowledge Base", "KBA 3623492 / KBA 2763512", "プロジェクト在庫 Q の PGI 不可（E2・6GD）／請求計画での請求不可", "userapps.support.sap.com（要ログイン）"),
    ("スコープアイテム", "6GD / IYT / 7DM", "Cloud の ETO（E2・E21）／MTO+VC／アドバンスド社内販売+MTO+VC", "help.sap.com/docs/s4hana-cloud-best-practices"),
]


def main():
    wb = Workbook()
    ws = wb.active; ws.title = "0_概要"
    kv(ws, "受注形態の横断比較 — 講義テキスト（概要）", "SAP S/4HANA On-Premise 基準 / サイト版 saporderflow/ と同期", SUMMARY, [22, 110])

    ws = wb.create_sheet("1_比較マトリクス")
    table(ws, "9 维度 横断比較マトリクス", "上から順に読むと「計画起点 → 库存区分 → 需求 → 製造/調達 → 成本 → 收钱」の因果が追える",
          ["# 维度", "MTS 見込生産", "MTO 受注生産", "ETO 受注設計生産", "ATO 受注組立", "VC バリアント設定"],
          MATRIX, [16, 30, 30, 34, 30, 34])

    ws = wb.create_sheet("2_早見表")
    table(ws, "早見表 1：戦略グループ → 所要量タイプ → 所要量クラス", "値は環境で異なる（特に 20/25/82/E2）。OPPJ/OPPS/OVZG と VA03 で実測する",
          ["戦略グループ", "名称", "所要量タイプ", "所要量クラス", "備考"], QUICK1, [16, 30, 16, 24, 52])
    r = 4 + len(QUICK1) + 2
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    c = ws.cell(r, 1, "早見表 2：在庫区分とテーブル"); c.font = F_HEAD; c.fill = FILL_HEAD
    for j, h in enumerate(["区分", "名称", "数量テーブル", "帰属先", "出荷の引当"], 1):
        cc = ws.cell(r + 1, j, h); cc.font = F_HEAD; cc.fill = FILL_HEAD; cc.alignment = CENTER; cc.border = BORDER
    for i, row in enumerate(QUICK2):
        for j, v in enumerate(row, 1):
            cc = ws.cell(r + 2 + i, j, v); cc.font = F_BODY; cc.alignment = WRAP; cc.border = BORDER
            if i % 2 == 1: cc.fill = FILL_ALT
    r2 = r + 2 + len(QUICK2) + 2
    ws.merge_cells(start_row=r2, start_column=1, end_row=r2, end_column=5)
    c = ws.cell(r2, 1, "早見表 3：明細カテゴリグループと業務の結び付き"); c.font = F_HEAD; c.fill = FILL_HEAD
    for j, h in enumerate(["グループ", "用途", "代表的な明細カテゴリ", "どこで効く", "—"], 1):
        cc = ws.cell(r2 + 1, j, h); cc.font = F_HEAD; cc.fill = FILL_HEAD; cc.alignment = CENTER; cc.border = BORDER
    for i, row in enumerate(QUICK3):
        for j, v in enumerate(list(row) + [""], 1):
            cc = ws.cell(r2 + 2 + i, j, v); cc.font = F_BODY; cc.alignment = WRAP; cc.border = BORDER
            if i % 2 == 1: cc.fill = FILL_ALT
    for j, w in enumerate([16, 30, 16, 24, 52], 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    ws = wb.create_sheet("3_E vs Q 対照")
    table(ws, "E（受注在庫）vs Q（プロジェクト在庫）— STEP 対照表", "同じ受注を 2 通りで走らせたときの一致点と相違点。実機手顺は sapmto/（MTO）と sapeto/（ETO）",
          ["STEP", "MTO（E）", "ETO（Q）", "同じ / 違うの本質"], EQ, [16, 46, 46, 46])

    ws = wb.create_sheet("4_演習解答")
    table(ws, "演習解答（形態の振り分け 5 件 / E vs Q 判別 8 問）", "講師用。受講者には配布せず、考えさせてから答え合わせに使う",
          ["#", "問い", "解答", "解説・採点ポイント"], SOLUTIONS, [12, 52, 34, 56])

    ws = wb.create_sheet("5_出典")
    table(ws, "出典・根拠", "比較表の値の根拠。具体値は必ず自システムで実測する",
          ["種別", "資料", "何の根拠に使ったか", "URL / 参照"], SOURCES, [20, 40, 52, 58])

    wb.save(OUT)
    print("saved:", OUT, "| sheets:", wb.sheetnames)


if __name__ == "__main__":
    main()
